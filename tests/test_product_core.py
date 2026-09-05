# -*- coding: utf-8 -*-
"""工程影像平台关键产品契约回归测试（标准库 unittest 可直接运行）。"""

from __future__ import annotations

import base64
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch


class DesktopUpdateCommandTest(unittest.TestCase):
    def test_update_targets_the_running_install_directory(self):
        from desktop_launcher import build_update_command

        temp_root = Path(tempfile.gettempdir())
        installer = temp_root / "album update.exe"
        install_dir = temp_root / "工程影像 客户安装目录"
        command = build_update_command(installer, install_dir)

        self.assertEqual(command[0], str(installer.resolve()))
        # 引号由 subprocess 自动添加：参数内嵌引号会被 Inno 当成路径的一部分，
        # 导致“文件夹名称不能包含下列任何字符”更新失败（v2.62 修复后契约）。
        self.assertIn("/DIR=" + str(install_dir.resolve()), command)
        self.assertFalse(any('"/DIR=' in arg for arg in command), "参数内不应内嵌引号")
        self.assertNotIn("/CLOSEAPPLICATIONS", command)


class VersionAssetSafetyTest(unittest.TestCase):
    def test_version_bump_never_cleans_assets_before_both_sources_exist(self):
        import bump_version

        with tempfile.TemporaryDirectory(prefix="album-version-assets-") as temp_dir:
            assets = Path(temp_dir)
            old_js = assets / "app-A20260822-2.64.js"
            old_gz = assets / "app-A20260822-2.64.js.gz"
            old_js.write_text("console.log('v1')", encoding="utf-8")
            old_gz.write_bytes(b"old-gzip-sentinel")
            with patch.object(bump_version, "ASSETS", assets):
                with self.assertRaisesRegex(SystemExit, "css 前端资源候选"):
                    bump_version._sync_asset_fingerprint("A20260825-2.65")
            self.assertTrue(old_js.exists(), "缺少 CSS 时不得提前重命名 JS")
            self.assertEqual(old_gz.read_bytes(), b"old-gzip-sentinel")

            old_css = assets / "app-A20260822-2.64.css"
            old_css.write_text("body{color:#123}", encoding="utf-8")
            with patch.object(bump_version, "ASSETS", assets):
                bump_version._sync_asset_fingerprint("A20260825-2.65")
            new_js = assets / "app-A20260825-2.65.js"
            new_css = assets / "app-A20260825-2.65.css"
            self.assertTrue(new_js.exists())
            self.assertTrue(new_css.exists())
            import gzip
            self.assertEqual(gzip.decompress((assets / f"{new_js.name}.gz").read_bytes()), new_js.read_bytes())
            self.assertEqual(gzip.decompress((assets / f"{new_css.name}.gz").read_bytes()), new_css.read_bytes())


class CloudAccountStateTest(unittest.TestCase):
    def test_missing_account_file_stays_logged_out(self):
        import cloud_client

        with tempfile.TemporaryDirectory(prefix="album-empty-account-") as temp_dir:
            missing = Path(temp_dir) / "cloud_account.json"
            with patch.object(cloud_client, "cloud_account_path", return_value=missing):
                self.assertEqual(cloud_client._read_cloud_account(), {})


class ServerProductContractTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="album-server-test-")
        os.environ["ALBUM_SERVER_DATA"] = cls.temp.name
        os.environ["ALBUM_SERVER_DB"] = str(Path(cls.temp.name) / "server.db")
        os.environ["ALBUM_ADMIN_USER"] = "test_admin"
        os.environ["ALBUM_ADMIN_PASS"] = "TestAdmin2026Secure"
        from server.app import create_app

        cls.app = create_app()
        cls.client = cls.app.test_client()

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def post(self, path, data, token=""):
        headers = {"Authorization": "Bearer " + token} if token else {}
        return self.client.post(path, json=data, headers=headers)

    def login(self, username, password):
        response = self.post("/api/login", {"username": username, "password": password})
        self.assertEqual(response.status_code, 200, response.get_json())
        return response.get_json()["token"]

    def test_login_and_me_return_same_stable_account_id(self):
        response = self.post(
            "/api/login",
            {"username": "test_admin", "password": "TestAdmin2026Secure"},
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        payload = response.get_json()
        self.assertIsInstance(payload["account"]["id"], int)
        # 新字段只能做加法，旧桌面端依赖的账号字段必须原样保留。
        self.assertTrue(
            {"username", "display_name", "role"}.issubset(payload["account"])
        )
        me = self.client.get(
            "/api/me",
            headers={"Authorization": "Bearer " + payload["token"]},
        )
        self.assertEqual(me.status_code, 200, me.get_json())
        me_payload = me.get_json()
        self.assertEqual(me_payload["account"]["id"], payload["account"]["id"])
        self.assertTrue(
            {"username", "display_name", "role"}.issubset(me_payload["account"])
        )
        self.assertIn("license", me_payload)

        # 同一账号重复登录时 subject 不随 token 或会话变化。
        second = self.post(
            "/api/login",
            {"username": "test_admin", "password": "TestAdmin2026Secure"},
        ).get_json()
        self.assertEqual(second["account"]["id"], payload["account"]["id"])

        # 模拟服务重启后重新建 Flask app，数据库里的稳定 ID 仍保持不变。
        from server.app import create_app

        restarted = create_app().test_client().post(
            "/api/login",
            json={"username": "test_admin", "password": "TestAdmin2026Secure"},
        )
        self.assertEqual(restarted.status_code, 200, restarted.get_json())
        self.assertEqual(
            restarted.get_json()["account"]["id"], payload["account"]["id"]
        )

    def test_product_config_invite_qualification_and_telemetry(self):
        product = self.client.get("/api/product-config").get_json()["product"]
        self.assertEqual(product["free"]["photo_limit"], 300)
        self.assertFalse(product["free"]["upload_expiry_enabled"])
        self.assertFalse(product["free"]["watermark_enabled"])
        self.assertEqual(product["pricing"]["price_label"], "原价")
        self.assertEqual(product["logging"]["retention_days"], 90)
        self.assertEqual(product["logging"]["max_records"], 5000)

        weak = self.post("/api/register", {"username": "weak", "password": "123456"})
        self.assertEqual(weak.status_code, 400)
        inviter = self.post(
            "/api/register",
            {"username": "inviter", "password": "Invite1234", "email": "a@example.com"},
        ).get_json()
        invite_code = inviter["account"]["invite_code"]
        inviter_token = self.login("inviter", "Invite1234")

        self.post(
            "/api/register",
            {"username": "invitee1", "password": "Invitee1234", "invite_code": invite_code},
        )
        first_token = self.login("invitee1", "Invitee1234")
        self.post("/api/activate", {"machine_id": "TEST-MACHINE-001"}, first_token)

        self.post(
            "/api/register",
            {"username": "invitee2", "password": "Invitee2234", "invite_code": invite_code},
        )
        second_token = self.login("invitee2", "Invitee2234")
        self.post("/api/activate", {"machine_id": "TEST-MACHINE-001"}, second_token)

        status = self.client.get(
            "/api/invite-status", headers={"Authorization": "Bearer " + inviter_token}
        ).get_json()
        self.assertEqual(status["registered_count"], 2)
        self.assertEqual(status["qualified_count"], 1)

        self.post(
            "/api/client-errors",
            {
                "source": "unittest",
                "message": "password=DoNotStore failed",
                "context": {
                    "token": "secret-token",
                    "photo_b64": "data:image/png;base64,AAAA",
                },
            },
        )
        admin_token = self.login("test_admin", "TestAdmin2026Secure")
        before_404 = self.client.get(
            "/api/admin/errors?resolved=0",
            headers={"Authorization": "Bearer " + admin_token},
        ).get_json()["summary"]["unresolved"]
        self.assertEqual(self.client.get("/definitely-not-a-real-route").status_code, 404)
        errors = self.client.get(
            "/api/admin/errors?resolved=0",
            headers={"Authorization": "Bearer " + admin_token},
        ).get_json()
        self.assertEqual(errors["summary"]["unresolved"], before_404)
        serialized = json.dumps(errors, ensure_ascii=False)
        self.assertNotIn("DoNotStore", serialized)
        self.assertNotIn("secret-token", serialized)
        self.assertNotIn("data:image", serialized)

        changed = self.post(
            "/api/admin/product-config",
            {"product": {"cloud_preview": {"thumb_max_side": 760}}},
            admin_token,
        ).get_json()["product"]
        self.assertEqual(changed["cloud_preview"]["thumb_max_side"], 760)
        self.assertEqual(changed["config_version"], 2)

        self.post(
            "/api/admin/version",
            {
                "version": "A20260809-9.99",
                "download_url": "https://example.invalid/update.exe",
                "sha256": "a" * 64,
                "size": 123456,
            },
            admin_token,
        )
        update = self.client.get("/api/check-update?current=A20260809-2.26").get_json()
        self.assertTrue(update["has_update"])
        self.assertTrue(update["mandatory"])
        self.assertEqual(update["sha256"], "a" * 64)
        self.assertEqual(update["size"], 123456)


class DesktopSecurityContractTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="album-desktop-test-")
        os.environ["ALBUM_DATA_DIR"] = cls.temp.name
        os.environ["ALBUM_TRIAL_FILE"] = str(Path(cls.temp.name) / "trial.json")
        os.environ["ALBUM_LICENSE_FILE"] = str(Path(cls.temp.name) / "license.dat")
        import app as desktop_app

        cls.desktop = desktop_app
        cls.client = desktop_app.app.test_client()

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_delete_password_is_hashed_and_cloud_thumb_is_recognizable(self):
        from PIL import Image

        with self.client.session_transaction() as session:
            session["username"] = "测试员"
        with patch.object(self.desktop.cloud_client, "report_error", return_value={"ok": True}):
            public = self.client.get("/api/project-info").get_json()
            self.assertFalse(public["project_info"]["delete_password_set"])
            self.assertNotIn("delete_password_hash", json.dumps(public))

            base = {"project_name": "测试工程", "project_code": "TEST-001"}
            self.assertEqual(self.client.post("/api/project-info", json=base).status_code, 400)
            self.assertEqual(
                self.client.post(
                    "/api/project-info", json={**base, "delete_password": "114405"}
                ).status_code,
                400,
            )
            password = "Delete7286"
            saved = self.client.post(
                "/api/project-info", json={**base, "delete_password": password}
            )
            self.assertEqual(saved.status_code, 200, saved.get_json())
            raw = Path(self.desktop.ENGINEERING_PROJECTS_PATH).read_text(encoding="utf-8")
            self.assertNotIn(password, raw)
            self.assertIn("delete_password_hash", raw)

        photos = self.desktop.manager._photos_dir("TEST-THUMB")
        photos.mkdir(parents=True, exist_ok=True)
        original = photos / "sample.jpg"
        Image.new("RGB", (2000, 1200), (42, 110, 160)).save(original, "JPEG", quality=92)
        with patch.object(self.desktop.cloud_client, "report_error", return_value={"ok": True}):
            raw_thumb = self.desktop._cloud_thumbnail_bytes(
                "TEST-THUMB",
                "sample.jpg",
                {"thumb_max_side": 720, "thumb_quality": 76, "max_thumb_b64": 600000},
            )
        self.assertLess(len(base64.b64encode(raw_thumb)), 600000)
        self.assertLessEqual(len(raw_thumb), 99 * 1024)
        with Image.open(io.BytesIO(raw_thumb)) as image:
            self.assertEqual(max(image.size), 720)

    def test_new_engineering_project_keeps_settings_and_other_images_isolated(self):
        """新工程必须不能继承或改写当前工程的规则、目录和照片容器。"""
        with self.client.session_transaction() as session:
            session["username"] = "工程隔离验收员"
        before = self.client.get("/api/project-info").get_json()
        previous_id = before["active_engineering_project_id"]
        previous_settings = self.client.get("/api/engineering-settings").get_json()["settings"]

        created = self.client.post(
            "/api/engineering-projects",
            json={
                "project_name": "专业隔离测试工程",
                "project_code": "ISOLATION-2026",
                "industry": "市政工程",
                "project_type": "道路工程",
                "delete_password": "Isolate7286",
            },
        )
        self.assertEqual(created.status_code, 200, created.get_json())
        payload = created.get_json()
        new_id = payload["active_engineering_project_id"]
        self.assertNotEqual(new_id, previous_id)
        self.assertTrue(Path(payload["backup"]).exists())
        self.assertEqual(payload["project_info"]["industry"], "市政工程")

        # 新工程的设置是默认副本；修改它不允许影响切回去的旧工程。
        self.assertEqual(
            self.client.post(
                "/api/engineering-settings",
                json={"settings": {"print_page_size": 2}},
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.post(
                "/api/other-images/nodes",
                json={"code": "QT-ISO", "name": "隔离目录"},
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.post("/api/engineering-projects/active", json={"id": previous_id}).status_code,
            200,
        )
        restored_settings = self.client.get("/api/engineering-settings").get_json()["settings"]
        self.assertEqual(restored_settings["print_page_size"], previous_settings["print_page_size"])
        restored_tree = self.client.get("/api/other-images/tree").get_json()["tree"]
        self.assertNotIn(
            "QT-ISO",
            [node.get("code") for node in restored_tree.get("nodes", [])],
        )

    def test_unit_photos_never_keep_a_process_number(self):
        """历史 GY 残留只能归到单元照片，真实工序照片仍使用 GX 编号。"""
        from PIL import Image
        from werkzeug.datastructures import FileStorage

        with self.client.session_transaction() as session:
            session["username"] = "编号修复验收员"

        project_code = "PROCESS-SCOPE-REPAIR"
        self.desktop.manager.create_project(
            project_code,
            "工序编号修复单元",
            source_fields={
                "单元工程名称": "工序编号修复单元",
                "单元工程编码": project_code,
                "划分类型": "划分工序单元",
                "工序名称列表": "混凝土浇筑",
            },
        )
        # 真实有工序的单元会在项目元数据中明确保存工序列表；这里模拟该现场。
        process_project = self.desktop.manager.get_project(project_code)
        process_project["processes"] = ["混凝土浇筑"]
        self.desktop.manager._write_metadata(project_code, process_project)

        def jpg_file(name):
            buffer = io.BytesIO()
            Image.new("RGB", (80, 60), (30, 120, 160)).save(buffer, "JPEG")
            buffer.seek(0)
            return FileStorage(stream=buffer, filename=name, content_type="image/jpeg")

        self.desktop.manager.add_photos(
            project_code, [jpg_file("construction.jpg")], category="construction"
        )
        self.desktop.manager.add_photos(
            project_code, [jpg_file("supervision.jpg")], category="supervision"
        )
        data = self.desktop.manager.get_project(project_code)
        for photo in data["photos"]:
            if photo["category"] == "construction":
                photo.update({"id": "SYGX064", "process_name": "", "process_code": "GY001"})
            else:
                photo.update({"id": "JYGX050", "process_name": "", "process_code": "GY001"})
        self.desktop.manager._write_metadata(project_code, data)

        repaired = self.client.get(f"/api/photos/{project_code}")
        self.assertEqual(repaired.status_code, 200, repaired.get_json())
        photos = repaired.get_json()["photos"]
        self.assertEqual(
            {photo["id"] for photo in photos if photo["category"] == "construction"},
            {"SYDY001"},
        )
        self.assertEqual(
            {photo["id"] for photo in photos if photo["category"] == "supervision"},
            {"JYDY001"},
        )
        self.assertTrue(all(not photo["process_code"] for photo in photos))

        actual_process = self.desktop.manager.add_photos(
            project_code,
            [jpg_file("actual-process.jpg")],
            category="construction",
            process_name="混凝土浇筑",
            process_code="GY999",  # 调用方给错也必须按工序名称纠正。
        )[0]
        self.assertEqual(actual_process["id"], "GY001-SYGX001")
        self.assertEqual(actual_process["process_code"], "GY001")

    def test_delete_photo_enqueues_explicit_uuid_without_rolling_back_local_delete(self):
        from PIL import Image
        from unittest.mock import MagicMock
        from werkzeug.datastructures import FileStorage
        from werkzeug.security import generate_password_hash

        with self.client.session_transaction() as session:
            session["username"] = "删除同步验收员"

        password = "DeleteRoute7286"
        engineering_project = dict(self.desktop.current_engineering_project())
        engineering_project["delete_password_hash"] = generate_password_hash(password)
        engineering_project["cloud_owner_id"] = "account-delete-route"

        def add_photo(project_code, filename):
            self.desktop.manager.create_project(
                project_code,
                f"{project_code} 删除同步验收单元",
                source_fields={"单元工程编码": project_code},
            )
            raw = io.BytesIO()
            Image.new("RGB", (320, 240), (35, 120, 105)).save(raw, "JPEG", quality=90)
            added = self.desktop.manager.add_photos(
                project_code,
                [FileStorage(stream=io.BytesIO(raw.getvalue()), filename=filename)],
            )
            stored = self.desktop.manager.get_project(project_code)["photos"][0]
            target = self.desktop.manager._photos_dir(project_code) / stored["filename"]
            return added[0], stored, target

        # 有永久 UUID：本地文件与元数据先删除，再以捕获到的稳定标识精确入队。
        project_code = "DELETE-UUID-ROUTE"
        added, stored, target = add_photo(project_code, "uuid-delete.jpg")
        photo_uuid = stored["photo_uuid"]
        service = MagicMock()
        service.enqueue_photo_delete.return_value = {"ok": True, "queued": True}
        with patch.object(
            self.desktop, "current_engineering_project", return_value=engineering_project
        ), patch.object(
            self.desktop, "desktop_mobile_sync_service", return_value=service
        ):
            response = self.client.post(
                f"/api/delete-photo/{project_code}/{added['id']}",
                json={"category": "construction", "safety_password": password},
            )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertTrue(response.get_json()["deleted"])
        service.enqueue_photo_delete.assert_called_once_with(
            photo_uuid,
            project_code,
            added["id"],
            project_id=str(engineering_project["id"]),
            cloud_account_id="account-delete-route",
        )
        self.assertFalse(target.exists())
        self.assertFalse(self.desktop.manager.get_project(project_code)["photos"])

        # 旧照片若还没有 UUID，只执行本地删除，绝不能凭目录或编号猜测云端对象。
        missing_code = "DELETE-NO-UUID-ROUTE"
        missing_added, missing_stored, missing_target = add_photo(
            missing_code, "no-uuid-delete.jpg"
        )
        with self.desktop.manager._project_write_lock(missing_code):
            data = self.desktop.manager.get_project(missing_code)
            data["photos"][0].pop("photo_uuid", None)
            self.desktop.manager._write_metadata(missing_code, data)
        missing_service = MagicMock()
        with patch.object(
            self.desktop, "current_engineering_project", return_value=engineering_project
        ), patch.object(
            self.desktop, "desktop_mobile_sync_service", return_value=missing_service
        ):
            missing_response = self.client.post(
                f"/api/delete-photo/{missing_code}/{missing_added['id']}",
                json={"category": "construction", "safety_password": password},
            )
        self.assertEqual(missing_response.status_code, 200, missing_response.get_json())
        self.assertTrue(missing_response.get_json()["deleted"])
        self.assertEqual(
            missing_response.get_json()["cloud_delete"]["code"], "photo_uuid_missing"
        )
        missing_service.enqueue_photo_delete.assert_not_called()
        self.assertFalse(missing_target.exists())
        self.assertFalse(self.desktop.manager.get_project(missing_code)["photos"])

        # 入队异常只会留在响应状态中；已经完成的本地删除不能因此被回滚。
        failure_code = "DELETE-QUEUE-FAIL-ROUTE"
        failure_added, failure_stored, failure_target = add_photo(
            failure_code, "queue-failure-delete.jpg"
        )
        failure_service = MagicMock()
        failure_service.enqueue_photo_delete.side_effect = RuntimeError("queue unavailable")
        with patch.object(
            self.desktop, "current_engineering_project", return_value=engineering_project
        ), patch.object(
            self.desktop, "desktop_mobile_sync_service", return_value=failure_service
        ):
            failure_response = self.client.post(
                f"/api/delete-photo/{failure_code}/{failure_added['id']}",
                json={"category": "construction", "safety_password": password},
            )
        self.assertEqual(failure_response.status_code, 200, failure_response.get_json())
        self.assertTrue(failure_response.get_json()["deleted"])
        self.assertEqual(
            failure_response.get_json()["cloud_delete"]["code"], "delete_queue_failed"
        )
        failure_service.enqueue_photo_delete.assert_called_once_with(
            failure_stored["photo_uuid"],
            failure_code,
            failure_added["id"],
            project_id=str(engineering_project["id"]),
            cloud_account_id="account-delete-route",
        )
        self.assertFalse(failure_target.exists())
        self.assertFalse(self.desktop.manager.get_project(failure_code)["photos"])

    def test_ocr_prefers_construction_content_and_free_export_has_no_watermark(self):
        lines = [
            "洺河龙泉缓洪工程",
            "施工区域：请输入内容…",
            "施 工 内 容：左岸连接堤",
            "拍摄时间：2026.06.05 15:34",
        ]
        fields = self.desktop.watermark_ocr.extract_fields(lines)
        self.assertEqual(fields["title"], "左岸连接堤")
        self.assertEqual(fields["construction_content"], "左岸连接堤")
        self.assertEqual(fields["construction_area"], "")
        self.assertEqual(fields["project_name"], "洺河龙泉缓洪工程")
        self.assertEqual(fields["photo_date"], "2026-06-05")
        self.assertEqual(self.desktop.watermark_text_for_user(), "")

        delegated_fields = {**fields, "title": "共享模块题名"}
        with patch.object(
            self.desktop.watermark_ocr._shared,
            "extract_fields",
            return_value=delegated_fields,
        ) as delegated_extract:
            self.assertIs(
                self.desktop.watermark_ocr.extract_fields(["测试水印"]),
                delegated_fields,
            )
            delegated_extract.assert_called_once_with(["测试水印"])

        with patch.object(
            self.desktop.mobile_recognition,
            "normalize_match_text",
            return_value="shared-normalized",
        ) as delegated_normalize:
            self.assertEqual(self.desktop.normalize_match_text("原始文本"), "shared-normalized")
            delegated_normalize.assert_called_once_with("原始文本")
        with patch.object(
            self.desktop.mobile_recognition,
            "match_signal_score",
            return_value=777,
        ) as delegated_score:
            self.assertEqual(
                self.desktop._match_signal_score("信号", "候选", "area", "unit"),
                777,
            )
            delegated_score.assert_called_once_with("信号", "候选", "area", "unit")

        shared_match = {
            "project_code": "YNLQ-SHARED",
            "project_name": "共享匹配单元",
            "confidence": "high",
            "confidence_label": "高",
            "ambiguous": False,
            "score": 980,
            "reasons": ["施工区域匹配单元“共享匹配单元”"],
            "alternatives": [{
                "unit_code": "YNLQ-SHARED",
                "project_code": "YNLQ-SHARED",
                "display_project_code": "YNLQ-SHARED",
                "project_name": "共享匹配单元",
                "unit_name": "共享单位工程",
                "part_name": "共享分部工程",
                "score": 980,
                "reasons": ["共享规则"],
            }],
            "candidates": ["不应泄露到旧返回契约"],
        }
        with patch.object(
            self.desktop.mobile_recognition,
            "match_photo",
            return_value=shared_match,
        ) as delegated_match, patch.object(
            self.desktop.manager,
            "list_projects",
            return_value=[{"project_code": "YNLQ-SHARED"}],
        ):
            compatible = self.desktop.match_photo_to_projects("现场.jpg", "现场/现场.jpg", {})
        delegated_match.assert_called_once()
        self.assertEqual(compatible["project_code"], "YNLQ-SHARED")
        self.assertEqual(set(compatible), {
            "project_code", "project_name", "confidence", "confidence_label",
            "ambiguous", "score", "reasons", "alternatives",
        })
        self.assertNotIn("unit_code", compatible["alternatives"][0])

    def test_multi_rule_match_process_disabled_and_print_presets(self):
        from PIL import Image
        from werkzeug.datastructures import FileStorage

        project = self.desktop.manager.create_project(
            "YNLQ-001-01-001",
            "左岸连接堤",
            source_fields={
                "合同工程名称": "洺河龙泉缓洪工程",
                "单位工程名称": "龙泉河单位工程",
                "分部工程名称": "左岸连接工程",
                "单元工程名称": "左岸连接堤",
                "单元工程编码": "YNLQ-001-01-001",
                "划分类型": "划分工序单元",
                "工序名称列表": "普通混凝土备仓、普通混凝土浇筑",
            },
        )
        self.assertEqual(project["division_type"], "unit")
        self.assertEqual(project["processes"], [])
        self.assertEqual(project["source_fields"]["工序名称列表"], "")

        result = self.desktop.match_photo_to_projects(
            "IMG_20260605.jpg",
            "批量照片/左岸连接堤/IMG_20260605.jpg",
            {
                "construction_area": "左岸连接堤",
                "construction_content": "混凝土浇筑完成",
                "project_name": "洺河龙泉缓洪工程",
            },
        )
        self.assertEqual(result["project_code"], project["project_code"])
        self.assertEqual(result["confidence"], "high")
        self.assertTrue(any("施工区域匹配单元" in reason for reason in result["reasons"]))

        image_bytes = io.BytesIO()
        Image.new("RGB", (640, 480), (100, 120, 140)).save(image_bytes, "JPEG")
        image_bytes.seek(0)
        self.desktop.manager.add_photos(
            project["project_code"],
            [FileStorage(stream=image_bytes, filename="施工照片.jpg")],
            category="construction",
            title="基槽验收",
            photo_date="2026-08-09",
        )
        copied = self.desktop.manager.copy_construction_to_supervision(project["project_code"])
        self.assertEqual(len(copied), 1)
        photos = self.desktop.manager.get_project(project["project_code"])["photos"]
        self.assertEqual(sum(photo["category"] == "construction" for photo in photos), 1)
        self.assertEqual(sum(photo["category"] == "supervision" for photo in photos), 1)

        with self.desktop.app.test_request_context("/?photo_preset=6inch"):
            settings = self.desktop.print_settings_from_request()
        self.assertEqual(settings["print_photo_preset"], "6inch")
        self.assertEqual(settings["print_photo_width_cm"], 15.2)
        self.assertEqual(settings["print_photo_height_cm"], 10.2)

    def test_remember_login_and_required_product_controls_are_visible(self):
        response = self.client.post(
            "/api/auth/login",
            json={"username": "长期登录测试员", "remember": True},
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        with self.client.session_transaction() as session:
            self.assertTrue(session.permanent)
            self.assertEqual(session.get("username"), "长期登录测试员")

        html = self.client.get("/").get_data(as_text=True)
        self.assertIn("复制施工照片", html)
        self.assertIn("4寸（10.2×7.6cm）", html)
        self.assertIn("审核批量导入结果", html)
        self.assertIn("默认打印", html)
        self.assertIn("每页照片数量", html)
        self.assertIn("文字标签位置", html)
        self.assertIn("printSetupMarginTop", html)

    def test_project_division_menu_exposes_excel_template_and_import_on_narrow_screens(self):
        from openpyxl import load_workbook

        html = self.client.get("/").get_data(as_text=True)
        self.assertIn('data-project-settings-tab="division"', html)
        self.assertIn('data-project-settings-panel="division"', html)
        self.assertIn("下载项目划分模板（Excel）", html)
        self.assertIn("选择项目划分文件导入", html)
        self.assertIn('aria-controls="projectImportMenu"', html)

        response = self.client.get("/api/project-template.xlsx")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data), read_only=True)
        self.assertEqual(workbook.active.title, "项目划分")
        headers = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
        self.assertIn("单元工程编码", headers)
        workbook.close()
        response.close()

        script_path = next((self.desktop.BASE_DIR / "assets").glob("app-A*.js"))
        script = script_path.read_text(encoding="utf-8")
        self.assertIn('menu?.classList.contains("settings-static-menu")', script)
        self.assertNotIn('$("#projectImportMenuBtn").addEventListener("click", (event) => {\n      event.stopPropagation();\n      $("#importFile").click();', script)

        styles_path = next((self.desktop.BASE_DIR / "assets").glob("app-A*.css"))
        styles = styles_path.read_text(encoding="utf-8")
        self.assertRegex(
            styles,
            r"(?s)@media \(max-width: 760px\).*?\.project-info-modal \.project-settings-tabs",
        )

    def test_ux_redesign_keeps_upload_primary_and_removes_metadata_exports(self):
        html = self.client.get("/").get_data(as_text=True)
        upload_pos = html.index('id="globalUploadBtn"')
        settings_pos = html.index('id="settingsMenuBtn"')
        self.assertLess(upload_pos, settings_pos, "上传按钮必须直接出现在顶部主菜单")
        self.assertIn('id="engineeringProjectSearch"', html)
        self.assertIn('list="engineeringProjectOptions"', html)
        self.assertIn('id="gateRegisterBtn" href="/license?register=1"', html)
        for tab in ("info", "division", "members", "upload", "print", "account", "device"):
            self.assertIn(f'data-project-settings-tab="{tab}"', html)
        self.assertIn('id="unitCategoryConstructionBtn"', html)
        self.assertIn('id="unitProcessTabs"', html)
        self.assertIn('id="photoSelectionBar"', html)
        self.assertIn('id="uploadFileQueue"', html)
        self.assertIn('class="upload-workspace"', html)
        self.assertIn('class="upload-target-column"', html)
        self.assertEqual(html.count('class="print-style-schematic"'), 6)
        self.assertIn('ps-paper ps-landscape ps-grid4', html)
        self.assertIn('ps-paper ps-portrait ps-stack2', html)
        self.assertIn('ps-paper ps-landscape ps-pair2', html)
        self.assertIn('id="addProcessModal"', html)
        self.assertNotIn('id="settingsExportCsvBtn"', html)
        self.assertNotIn('id="settingsExportJsonBtn"', html)

        routes = {rule.rule for rule in self.desktop.app.url_map.iter_rules()}
        self.assertNotIn("/api/export/<path:project_code>/metadata.csv", routes)
        self.assertNotIn("/api/export/<path:project_code>/metadata.json", routes)
        self.assertIn("/api/export/<path:project_code>/zip", routes)
        self.assertIn("/api/export/photo-statistics.xlsx", routes)
        self.assertIn("/api/engineering-package/export", routes)

        script_path = next((self.desktop.BASE_DIR / "assets").glob("app-A*.js"))
        script = script_path.read_text(encoding="utf-8")
        self.assertIn('form.append("process_name", processName)', script)
        self.assertIn('state.uploadQueuedFiles', script)
        self.assertIn('class="recent-upload-thumb', script)
        self.assertIn('rec.thumb_url', script)
        self.assertIn('modal.classList.toggle("is-current-unit-upload", currentUnitOnly)', script)
        self.assertNotIn('$("#gateRegisterBtn").addEventListener', script)
        self.assertNotIn('if (manual || !state.auth.user) renderManualUpdate(data)', script)
        self.assertIn('versionBadge?.classList.toggle("update-available", Boolean(hasUpdate))', script)
        self.assertIn('// 新版本对登录前后都必须明显可见', script)
        self.assertIn('openUploadChoiceModal(normalizePhotoCategory(category), processName, true)', script)
        self.assertIn('class="photo-selection-state"', script)
        self.assertIn('syncPhotoCardSelection(card, next)', script)
        self.assertNotIn('data-print-photo=', script)
        self.assertNotIn("metadata.csv", script)
        self.assertNotIn("metadata.json", script)

        styles_path = next((self.desktop.BASE_DIR / "assets").glob("app-A*.css"))
        styles = styles_path.read_text(encoding="utf-8")
        self.assertIn(".recent-upload-thumb img", styles)
        self.assertIn("--ux-hover-soft", styles)
        self.assertIn('button[type="submit"]:not(.secondary)', styles)
        self.assertIn("#uploadChoiceModal.is-current-unit-upload .upload-target-column", styles)
        self.assertIn(".print-style-schematic", styles)
        self.assertIn(".print-mode-card.print-style-card:not(:disabled):not(.active):hover", styles)
        self.assertIn("#unitView .photo-selection-state", styles)
        self.assertIn(".auth-register-link", styles)
        self.assertIn('#checkUpdateBtn.update-available', styles)
        self.assertIn('.app-version-badge.update-available::after', styles)
        self.assertIn('.update-overlay[data-mode="manual"] .update-actions .button-link', styles)
        self.assertIn('content: " · 可更新"', styles)
        self.assertRegex(
            styles,
            r"\.guide-overlay\s*\{[^}]*pointer-events:\s*none;",
            "尚未激活的工作台引导遮罩不能阻挡登录和注册按钮",
        )
        self.assertRegex(
            styles,
            r"\.guide-overlay\.active\s*\{[^}]*pointer-events:\s*auto;",
            "只允许已激活的新手引导接管鼠标操作",
        )
        self.assertRegex(
            styles,
            r"(?s)@media \(max-width: 760px\).*?\.overview-grid > \.overview-board\s*\{.*?grid-column: 1 !important;.*?grid-row: auto !important;",
        )

        license_html = self.client.get("/license?register=1").get_data(as_text=True)
        self.assertIn('let registrationRequested = new URLSearchParams(location.search).get("register") === "1"', license_html)
        self.assertIn("if (registrationRequested) showRegistrationForm();", license_html)

    def test_recent_uploads_return_thumbnail_urls_for_frontend_cards(self):
        from PIL import Image
        from werkzeug.datastructures import FileStorage

        with self.client.session_transaction() as session:
            session["username"] = "缩略图验收员"
        project = self.desktop.manager.create_project(
            "RECENT-THUMB-001",
            "最近上传缩略图验收单元",
            source_fields={
                "合同工程名称": "缩略图验收合同",
                "单位工程名称": "缩略图验收单位工程",
                "分部工程名称": "缩略图验收分部",
                "单元工程名称": "最近上传缩略图验收单元",
                "单元工程编码": "RECENT-THUMB-001",
            },
        )
        raw = io.BytesIO()
        Image.new("RGB", (640, 480), (36, 112, 105)).save(raw, "JPEG")
        raw.seek(0)
        self.desktop.manager.add_photos(
            project["project_code"],
            [FileStorage(stream=raw, filename="最近上传缩略图.jpg")],
            title="最近上传缩略图",
            photo_date="2026-09-01",
            uploaded_by="缩略图验收员",
        )

        response = self.client.get("/api/recent-uploads?limit=100&scope=construction")
        self.assertEqual(response.status_code, 200, response.get_json())
        records = response.get_json()["records"]
        record = next(item for item in records if item["project_code"] == project["project_code"])
        self.assertTrue(record["thumb_url"].startswith(f"/thumbs/{project['project_code']}/"))
        thumbnail = self.client.get(record["thumb_url"])
        self.assertEqual(thumbnail.status_code, 200)
        self.assertEqual(thumbnail.mimetype, "image/webp")
        thumbnail.close()

    def test_print_defaults_show_complete_photo_and_readable_photographer_caption(self):
        with self.desktop.app.test_request_context("/?style=default"):
            settings = self.desktop.print_settings_from_request()
        self.assertEqual(settings["print_image_fit"], "contain")
        self.assertEqual(settings["print_caption_font_size"], 10.0)
        self.assertEqual(settings["print_caption_size_mm"], 7.5)
        self.assertIn("photographer", settings["print_caption_fields"])
        self.assertIn("part_name", settings["print_caption_fields"])
        legacy_print_fields = self.desktop._normalize_engineering_settings({
            "print_caption_fields": ["title", "id", "photo_date"],
        })
        self.assertEqual(
            legacy_print_fields["print_caption_fields"],
            ["title", "id", "part_name", "photo_date"],
        )
        explicit_fields = self.desktop._normalize_engineering_settings({
            "print_caption_fields": ["title", "id", "photo_date"],
            "print_caption_fields_schema_version": 2,
        })
        self.assertNotIn("part_name", explicit_fields["print_caption_fields"])

        with self.desktop.app.test_request_context("/?style=custom&image_fit=cover"):
            custom = self.desktop.print_settings_from_request()
        self.assertEqual(custom["print_image_fit"], "cover")

        project = {
            "project_code": "PRINT-CHECK",
            "project_name": "打印验收单元",
            "source_fields": {"分部工程名称": "左岸连接工程"},
        }
        photos = self.desktop.enrich_print_photos(project, [{
            "id": "PRINT-CHECK-SY001",
            "title": "四边完整验收照片",
            "photo_date": "2026-08-28",
            "photographer": "张三",
            "url": "/photos/PRINT-CHECK/sample.jpg",
            "orientation": "portrait",
            "rotation": 0,
        }], settings)
        pages = self.desktop.layout_fixed_photo_pages(
            photos,
            layout="landscape",
            page_size=4,
            photo_width_mm=127,
            photo_height_mm=89,
            print_settings=settings,
        )
        self.assertEqual(pages[0][0]["_rotate"], 90)
        with self.desktop.app.test_request_context("/"):
            rendered = self.desktop.render_template(
                "print.html",
                project=project,
                photos=photos,
                project_groups=[{"level": "", "pages": [{"unit_name": "打印验收单元", "items": pages[0]}]}],
                category="construction",
                category_label="施工单元影像",
                print_layout="landscape",
                print_page_size=4,
                print_settings=settings,
                watermark_text="",
            )
        self.assertIn("object-fit: contain", rendered)
        self.assertIn("font-size: 10.0px", rendered)
        self.assertIn('class="pr-90"', rendered)
        self.assertIn("分部工程：左岸连接工程", rendered)
        self.assertIn("摄影者：张三", rendered)

    def test_single_photo_print_is_centered_and_keeps_selected_photo_size(self):
        settings = {
            "print_margin_top_mm": 6,
            "print_margin_right_mm": 6,
            "print_margin_bottom_mm": 6,
            "print_margin_left_mm": 6,
            "print_horizontal_gap_mm": 2,
            "print_vertical_gap_mm": 2,
            "print_caption_position": "bottom",
            "print_caption_size_mm": 7.5,
            "print_show_page_title": True,
            "print_single_alignment": "center",
        }
        pages = self.desktop.layout_fixed_photo_pages(
            [{"id": "SY-001", "detail_rows": [{"label": "题名", "value": "单张居中"}]}],
            layout="landscape",
            page_size=1,
            photo_width_mm=102,
            photo_height_mm=76,
            print_settings=settings,
        )
        item = pages[0][0]
        printable_width = 297 - settings["print_margin_left_mm"] - settings["print_margin_right_mm"]
        printable_height = 210 - settings["print_margin_top_mm"] - settings["print_margin_bottom_mm"] - 7
        self.assertAlmostEqual(item["_left"], (printable_width - item["_w"]) / 2, places=1)
        self.assertAlmostEqual(item["_top"], (printable_height - item["_block_h"]) / 2, places=1)

        with self.desktop.app.test_request_context("/?style=custom&photo_preset=6inch&single_alignment=top"):
            configured = self.desktop.print_settings_from_request()
        self.assertEqual(configured["print_photo_width_cm"], 15.2)
        self.assertEqual(configured["print_photo_height_cm"], 10.2)
        self.assertEqual(configured["print_single_alignment"], "top")

    def test_last_print_replays_exact_photo_ids_after_new_photos_are_added(self):
        from PIL import Image
        from werkzeug.datastructures import FileStorage

        project = self.desktop.manager.create_project(
            "LAST-PRINT-EXACT",
            "上次打印精确照片验收",
            source_fields={"分部工程名称": "左岸分部工程"},
        )

        def add_photo(title):
            source = io.BytesIO()
            Image.new("RGB", (640, 480), (30, 120, 150)).save(source, "JPEG")
            source.seek(0)
            return self.desktop.manager.add_photos(
                project["project_code"],
                [FileStorage(stream=source, filename=f"{title}.jpg")],
                category="construction",
                title=title,
                photo_date="2026-08-29",
            )[0]

        first = add_photo("上次打印照片一")
        second = add_photo("上次打印照片二")
        initial = self.client.get(
            f"/print/construction/{project['project_code']}?ids={first['id']},{second['id']}"
        )
        self.assertEqual(initial.status_code, 200, initial.get_data(as_text=True))

        saved_job = self.desktop.read_engineering_settings()["last_print_job"]
        self.assertEqual(saved_job["photo_count"], 2)
        self.assertEqual(
            saved_job["groups"][0]["items"][0]["photo_ids"],
            [first["id"], second["id"]],
        )

        add_photo("后来新增不能混入")
        replay = self.client.get("/print/last")
        rendered = replay.get_data(as_text=True)
        self.assertEqual(replay.status_code, 200, rendered)
        self.assertIn("上次打印照片一", rendered)
        self.assertIn("上次打印照片二", rendered)
        self.assertNotIn("后来新增不能混入", rendered)

        normalized = self.desktop._normalize_engineering_settings({
            "last_print_job": {
                "groups": [{
                    "level": "单元照片",
                    "items": [{
                        "project_code": "LAST-PRINT-EXACT",
                        "category": "construction",
                        "photo_ids": [first["id"], first["id"], ""],
                    }, {
                        "project_code": "DROP",
                        "category": "unknown",
                        "photo_ids": ["DROP-SY001"],
                    }],
                }],
            },
        })
        self.assertEqual(normalized["last_print_job"]["photo_count"], 1)
        self.assertEqual(
            normalized["last_print_job"]["groups"][0]["items"][0]["photo_ids"],
            [first["id"]],
        )

    def test_print_layout_keeps_captions_readable_and_page_numbers_unique(self):
        settings = {
            "print_margin_top_mm": 6,
            "print_margin_right_mm": 6,
            "print_margin_bottom_mm": 6,
            "print_margin_left_mm": 6,
            "print_horizontal_gap_mm": 2,
            "print_vertical_gap_mm": 2,
            "print_caption_position": "bottom",
            "print_caption_size_mm": 7.5,
            "print_show_page_title": True,
        }
        detailed_photo = {
            "id": "SY-DETAIL",
            "detail_rows": [
                {"label": "题名", "value": "完整显示文字"},
                {"label": "编号", "value": "SY-DETAIL"},
                {"label": "拍摄日期", "value": "2026-08-29"},
                {"label": "摄影人", "value": "测试员"},
            ],
        }
        pages = self.desktop.layout_fixed_photo_pages(
            [
                detailed_photo,
                dict(detailed_photo, id="SY-DETAIL-2"),
                dict(detailed_photo, id="SY-DETAIL-3"),
                dict(detailed_photo, id="SY-DETAIL-4"),
            ],
            layout="landscape",
            page_size=4,
            photo_width_mm=127,
            photo_height_mm=89,
            print_settings=settings,
        )
        self.assertEqual(len(pages), 2, "说明区不足时应自动分页，不能把文字压进照片")
        self.assertGreaterEqual(pages[0][0]["_caption_h"], 11.0)

        groups = [
            {"project_code": "U-1", "project_name": "一号单元", "level": "单元照片", "pages": [pages[0], pages[1]]},
            {"project_code": "U-2", "project_name": "二号单元", "level": "工序照片", "pages": [[detailed_photo]]},
        ]
        numbered = self.desktop.assign_unique_print_page_numbers(groups)
        all_pages = [page for group in numbered for page in group["pages"]]
        self.assertEqual([page["print_page_number"] for page in all_pages], [1, 2, 3])
        self.assertEqual([page["unit_page_number"] for page in all_pages], [1, 2, 1])

        normalized = self.desktop._normalize_engineering_settings({
            "last_print_settings": {"page_size": "1", "photo_preset": "6inch", "unknown": "discard"},
        })
        self.assertEqual(normalized["last_print_settings"], {"page_size": "1", "photo_preset": "6inch"})

    def test_editing_photo_without_orientation_keeps_detected_portrait(self):
        from PIL import Image
        from werkzeug.datastructures import FileStorage

        project = self.desktop.manager.create_project(
            "PORTRAIT-EDIT-CHECK",
            "竖图编辑验收单元",
            source_fields={"单元工程编码": "PORTRAIT-EDIT-CHECK"},
        )
        source = io.BytesIO()
        Image.new("RGB", (800, 1200), (80, 130, 170)).save(source, "JPEG", quality=90)
        source.seek(0)
        added = self.desktop.manager.add_photos(
            project["project_code"],
            [FileStorage(stream=source, filename="竖图.jpg")],
            title="竖图",
            photo_date="2026-08-28",
        )
        self.assertEqual(added[0]["orientation"], "portrait")

        with self.client.session_transaction() as session:
            session["username"] = "方向测试员"
        response = self.client.patch(
            f"/api/photos/{project['project_code']}/{added[0]['id']}",
            json={
                "title": "只改题名",
                "photo_date": "2026-08-28",
                "photographer": "方向测试员",
                "description": "",
                "process_name": "",
            },
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["photo"]["orientation"], "portrait")

    def test_photo_cards_show_readable_date_and_photographer_or_uploader(self):
        script_path = next((self.desktop.BASE_DIR / "assets").glob("app-A*.js"))
        styles_path = next((self.desktop.BASE_DIR / "assets").glob("app-A*.css"))
        script = script_path.read_text(encoding="utf-8")
        styles = styles_path.read_text(encoding="utf-8")

        self.assertIn('const person = photographer', script)
        self.assertIn('`上传者 ${uploader}`', script)
        self.assertIn('class="photo-card-details"', script)
        self.assertIn(".photo-card-details", styles)
        self.assertNotIn(".photo-card-title { font-size: 7.5px; }", styles)
        self.assertNotIn(".photo-card-title { font-size: 9px; }", styles)

    def test_default_print_is_four_5inch_photos_and_upload_is_optimized(self):
        import hashlib
        import uuid

        from PIL import Image
        from werkzeug.datastructures import FileStorage

        photos = [
            {"id": f"SY-{index}", "detail_rows": [{"label": "题名", "value": f"照片{index}"}]}
            for index in range(1, 5)
        ]
        settings = {
            "print_margin_top_mm": 4,
            "print_margin_right_mm": 4,
            "print_margin_bottom_mm": 4,
            "print_margin_left_mm": 4,
            "print_horizontal_gap_mm": 2,
            "print_vertical_gap_mm": 2,
            "print_caption_position": "bottom",
            "print_caption_size_mm": 7.5,
            "print_show_page_title": True,
        }
        pages = self.desktop.layout_fixed_photo_pages(
            photos,
            layout="landscape",
            page_size=4,
            photo_width_mm=127,
            photo_height_mm=89,
            print_settings=settings,
        )
        self.assertEqual(len(pages), 1)
        self.assertEqual(len(pages[0]), 4)
        self.assertTrue(all(item["_caption_position"] == "bottom" for item in pages[0]))
        printable_height = 210 - settings["print_margin_top_mm"] - settings["print_margin_bottom_mm"] - 7
        self.assertTrue(all(item["_top"] + item["_block_h"] <= printable_height for item in pages[0]))

        source = io.BytesIO()
        Image.effect_noise((4200, 2800), 48).convert("RGB").save(source, "JPEG", quality=100)
        source_bytes = source.getvalue()
        source.seek(0)
        target = self.desktop.manager._photos_dir("OPT-CHECK") / "optimized-upload.jpg"
        result = self.desktop.manager._save_optimized_upload(
            FileStorage(stream=source, filename="现场原图.jpg"), target
        )
        self.assertTrue(result["optimized"])
        self.assertLess(target.stat().st_size, len(source_bytes))
        self.assertLessEqual(target.stat().st_size, 300 * 1024)
        backup = self.desktop.manager._original_path("OPT-CHECK", result["original_backup"])
        self.assertTrue(backup.exists())
        self.assertEqual(backup.stat().st_size, len(source_bytes))
        self.assertTrue(self.desktop.manager._ensure_thumbnail("OPT-CHECK", target.name))
        self.assertTrue((target.parent / self.desktop.manager._thumb_filename(target.name)).exists())
        with Image.open(target) as optimized:
            self.assertLessEqual(max(optimized.size), 3200)

        # 手机上传改为服务器短期中转原图：即使文件明显大于 200KB，
        # 桌面端也必须原字节归档，后续优化入口不得再次编码。
        mobile_project = self.desktop.manager.create_project(
            "MOBILE-ORIGINAL-CHECK",
            "手机原图验收单元",
            source_fields={"单元工程编码": "MOBILE-ORIGINAL-CHECK"},
        )
        mobile_source = io.BytesIO()
        Image.effect_noise((1800, 1200), 72).convert("RGB").save(
            mobile_source, "JPEG", quality=98
        )
        mobile_raw = mobile_source.getvalue()
        self.assertGreater(len(mobile_raw), 200 * 1024)
        mobile_uuid = str(uuid.uuid4())
        mobile_name = f"mobile-{mobile_uuid}.jpg"
        added = self.desktop.manager.add_photos(
            mobile_project["project_code"],
            [FileStorage(stream=io.BytesIO(mobile_raw), filename=mobile_name)],
            category="construction",
            per_file_meta={mobile_name: {
                "photo_uuid": mobile_uuid,
                "preserve_mobile_original": True,
                "mobile_original_received": True,
                "source_original_name": "IMG_20260825_090000.jpg",
                "sha256": hashlib.sha256(mobile_raw).hexdigest(),
                "pixel_width": 1800,
                "pixel_height": 1200,
                "photographer_account_id": "member-8",
                "uploaded_by_account_id": "member-8",
                "project_owner_account_id": "owner-1",
                "capture_source": "camera",
            }},
        )
        saved = self.desktop.manager.get_project(mobile_project["project_code"])["photos"][0]
        mobile_target = (
            self.desktop.manager._photos_dir(mobile_project["project_code"])
            / saved["filename"]
        )
        self.assertEqual(mobile_target.read_bytes(), mobile_raw)
        self.assertFalse(saved["optimized"])
        self.assertTrue(saved["mobile_original_received"])
        self.assertEqual(saved["photographer_account_id"], "member-8")
        self.assertEqual(saved["uploaded_by_account_id"], "member-8")
        self.assertEqual(saved["project_owner_account_id"], "owner-1")
        self.assertEqual(saved["capture_source"], "camera")
        self.desktop.manager.optimize_existing_photo(
            mobile_project["project_code"], added[0]["id"]
        )
        self.assertEqual(mobile_target.read_bytes(), mobile_raw)
        self.assertEqual(
            hashlib.sha256(mobile_target.read_bytes()).hexdigest(),
            hashlib.sha256(mobile_raw).hexdigest(),
        )
        mobile_target.write_bytes(b"tampered-local-original")
        repaired = self.desktop.manager.repair_mobile_original(
            mobile_project["project_code"],
            mobile_uuid,
            mobile_raw,
            {
                "sha256": hashlib.sha256(mobile_raw).hexdigest(),
                "size_bytes": len(mobile_raw),
                "mime_type": "image/jpeg",
                "extension": ".jpg",
                "width": 1800,
                "height": 1200,
            },
            "IMG_20260825_090000.jpg",
        )
        self.assertEqual(mobile_target.read_bytes(), mobile_raw)
        self.assertEqual(repaired["sha256"], hashlib.sha256(mobile_raw).hexdigest())
        repair_backup = (
            self.desktop.manager._project_dir(mobile_project["project_code"])
            / "_mobile_sync_backups"
            / repaired["mobile_repair_backup"]
        )
        self.assertEqual(repair_backup.read_bytes(), b"tampered-local-original")

        png_project = self.desktop.manager.create_project(
            "MOBILE-PNG-CHECK",
            "手机 PNG 原图验收单元",
            source_fields={"单元工程编码": "MOBILE-PNG-CHECK"},
        )
        png_source = io.BytesIO()
        Image.effect_noise((900, 700), 68).convert("RGB").save(png_source, "PNG")
        png_raw = png_source.getvalue()
        png_uuid = str(uuid.uuid4())
        png_name = f"mobile-{png_uuid}.png"
        self.desktop.manager.add_photos(
            png_project["project_code"],
            [FileStorage(stream=io.BytesIO(png_raw), filename=png_name)],
            per_file_meta={png_name: {
                "photo_uuid": png_uuid,
                "preserve_mobile_original": True,
                "mobile_original_received": True,
                "sha256": hashlib.sha256(png_raw).hexdigest(),
                "pixel_width": 900,
                "pixel_height": 700,
            }},
        )
        png_saved = self.desktop.manager.get_project(png_project["project_code"])["photos"][0]
        png_target = self.desktop.manager._photos_dir(png_project["project_code"]) / png_saved["filename"]
        self.assertEqual(png_target.read_bytes(), png_raw)
        self.assertEqual(
            hashlib.sha256(png_target.read_bytes()).hexdigest(),
            hashlib.sha256(png_raw).hexdigest(),
        )

        failure_project = self.desktop.manager.create_project(
            "MOBILE-WRITE-FAIL-CHECK",
            "手机原图失败清理单元",
            source_fields={"单元工程编码": "MOBILE-WRITE-FAIL-CHECK"},
        )
        failure_name = f"mobile-{uuid.uuid4()}.jpg"
        with patch.object(
            self.desktop.manager, "_write_metadata", side_effect=OSError("metadata full")
        ):
            with self.assertRaises(OSError):
                self.desktop.manager.add_photos(
                    failure_project["project_code"],
                    [FileStorage(stream=io.BytesIO(mobile_raw), filename=failure_name)],
                    per_file_meta={failure_name: {
                        "photo_uuid": str(uuid.uuid4()),
                        "preserve_mobile_original": True,
                        "mobile_original_received": True,
                    }},
                )
        failure_dir = self.desktop.manager._photos_dir(failure_project["project_code"])
        self.assertFalse(list(failure_dir.glob("*.mobile.tmp")))
        self.assertFalse([
            path for path in failure_dir.iterdir()
            if path.suffix.lower() in {".jpg", ".jpeg", ".png"}
        ])

        replace_target = failure_dir / "replace-failure.jpg"
        with patch.object(self.desktop.os, "replace", side_effect=OSError("replace blocked")):
            with self.assertRaises(OSError):
                self.desktop.manager._save_exact_mobile_original(
                    FileStorage(stream=io.BytesIO(mobile_raw), filename="replace-failure.jpg"),
                    replace_target,
                )
        self.assertFalse(replace_target.exists())
        self.assertFalse(replace_target.with_suffix(".jpg.mobile.tmp").exists())

    def test_bootstrap_is_compact_paginated_and_assets_are_cached(self):
        with self.client.session_transaction() as session:
            session["username"] = "性能验收用户"
        for index in range(1, 26):
            self.desktop.manager.create_project(
                f"PERF-{index:03d}",
                f"性能单元{index:03d}",
                source_fields={
                    "合同工程名称": "性能测试合同",
                    "单位工程名称": "性能测试单位工程",
                    "单位工程编码": "PERF-U",
                    "分部工程名称": "性能测试分部",
                    "分部工程编码": "PERF-P",
                    "单元工程名称": f"性能单元{index:03d}",
                    "单元工程编码": f"PERF-{index:03d}",
                },
                project_order=index,
            )
        response = self.client.get("/api/bootstrap")
        self.assertEqual(response.status_code, 200)
        self.assertLess(len(response.data), 150 * 1024)
        data = response.get_json()
        parts = [
            part
            for contract in data["hierarchy"]
            for unit in contract["units"]
            for part in unit["parts"]
            if part["name"] == "性能测试分部"
        ]
        self.assertEqual(len(parts), 1)
        units = self.client.get(
            f"/api/units?part_key={parts[0]['part_key']}&page=1&page_size=20"
        ).get_json()
        self.assertEqual(units["total"], 25)
        self.assertEqual(len(units["projects"]), 20)
        self.assertEqual(units["pages"], 2)

        index_response = self.client.get("/", headers={"Accept-Encoding": "gzip"})
        self.assertEqual(index_response.headers.get("Content-Encoding"), "gzip")
        # 从 index.html 提取版本化 asset 文件名，避免 bump 版本后硬编码失效
        import gzip as _gzip
        import re as _re
        index_html = _gzip.decompress(index_response.data).decode("utf-8")
        self.assertIn('data-project-settings-tab="members"', index_html)
        self.assertIn('id="projectMembersList"', index_html)
        self.assertIn('id="createProjectInviteBtn"', index_html)
        asset_match = _re.search(r"/assets/(app-A\d{8}-\d+\.\d+\.js)", index_html)
        self.assertIsNotNone(asset_match, "index.html 中未找到版本化 JS asset")
        asset_name = asset_match.group(1)
        asset_response = self.client.get(
            f"/assets/{asset_name}", headers={"Accept-Encoding": "gzip"}
        )
        self.assertEqual(asset_response.headers.get("Content-Encoding"), "gzip")
        self.assertIn("immutable", asset_response.headers.get("Cache-Control", ""))
        asset_body = _gzip.decompress(asset_response.data)
        self.assertEqual(
            asset_body,
            (self.desktop.BASE_DIR / "assets" / asset_name).read_bytes(),
            "预压缩 JS 必须与当前源码完全一致，不能把旧界面缓存一年",
        )
        asset_text = asset_body.decode("utf-8")
        self.assertIn("/api/cloud/project-ownership", asset_text)
        self.assertIn("/api/cloud/wechat-binding", asset_text)
        self.assertIn("/api/cloud/project-members", asset_text)
        self.assertIn("/api/cloud/project-invites", asset_text)
        self.assertIn("已阻止串号", asset_text)
        asset_response.close()
        index_response.close()

    def test_z_cloud_project_ownership_is_backed_up_and_cannot_cross_accounts(self):
        project = self.desktop.default_engineering_project()
        project.update({
            "id": "ownership-test",
            "project_code": "OWN-001",
            "project_name": "账号归属测试工程",
            "delete_password_hash": self.desktop.generate_password_hash("OwnSafe2026"),
            "cloud_owner_id": "",
            "cloud_owner_username": "",
            "cloud_owner_claimed_at": "",
        })
        self.desktop.write_engineering_projects({
            "active_id": project["id"],
            "projects": [project],
        })

        with patch.object(
            self.desktop.cloud_client, "desktop_sync_account_id", return_value="account-A"
        ), patch.object(
            self.desktop.cloud_client,
            "desktop_sync_account_username",
            return_value="tenant_a",
        ):
            claimed = self.desktop.claim_current_project_cloud_ownership()
            self.assertEqual(claimed["state"], "matched")
            self.assertEqual(claimed["project_ref"], "account-A:ownership-test")
            backup_dir = Path(claimed["backup_dir"])
            self.assertTrue((backup_dir / self.desktop.ENGINEERING_PROJECTS_PATH.name).exists())
            self.assertTrue((backup_dir / self.desktop.PROJECT_INFO_PATH.name).exists())

            with self.client.session_transaction() as session:
                session["username"] = "tenant_a"
            saved = self.client.post(
                "/api/project-info",
                json={
                    "id": "ownership-test",
                    "project_code": "OWN-001",
                    "project_name": "账号归属测试工程（编辑后）",
                },
            )
            self.assertEqual(saved.status_code, 200, saved.get_json())
            preserved = self.desktop.current_engineering_project()
            self.assertEqual(preserved["cloud_owner_id"], "account-A")
            self.assertEqual(preserved["cloud_owner_username"], "tenant_a")
            self.assertEqual(preserved["cloud_project_ref"], "account-A:ownership-test")

        with patch.object(
            self.desktop.cloud_client, "desktop_sync_account_id", return_value="account-B"
        ), patch.object(
            self.desktop.cloud_client,
            "desktop_sync_account_username",
            return_value="tenant_b",
        ):
            status = self.desktop.current_project_cloud_ownership()
            self.assertEqual(status["state"], "mismatch")
            self.assertFalse(status["sync_allowed"])
            with self.assertRaisesRegex(ValueError, "不能直接改绑"):
                self.desktop.claim_current_project_cloud_ownership()

    def test_z_project_member_proxy_uses_owner_local_project_and_archive_machine(self):
        with self.client.session_transaction() as session:
            session["username"] = "project-owner"
        ownership = {
            "state": "matched",
            "local_project_id": "owner-local-project",
            "project_ref": "42:owner-local-project",
        }
        members_payload = {
            "ok": True,
            "members": [{"account_id": 42, "role": "owner"}],
        }
        with patch.object(
            self.desktop, "current_project_cloud_ownership", return_value=ownership
        ), patch.object(
            self.desktop.licensing, "machine_id", return_value="ARCHIVE-PC"
        ), patch.object(
            self.desktop.cloud_client,
            "desktop_project_members",
            return_value=members_payload,
        ) as get_members, patch.object(
            self.desktop.cloud_client,
            "desktop_create_project_invite",
            return_value={"ok": True, "invite": {"invite_id": 9, "code": "ABCDEF0123456789"}},
        ) as create_invite, patch.object(
            self.desktop.cloud_client,
            "desktop_update_project_member",
            return_value={"ok": True},
        ) as update_member, patch.object(
            self.desktop.cloud_client,
            "desktop_remove_project_member",
            return_value={"ok": True},
        ) as remove_member, patch.object(
            self.desktop.cloud_client,
            "desktop_revoke_project_invite",
            return_value={"ok": True},
        ) as revoke_invite:
            members = self.client.get("/api/cloud/project-members")
            invite = self.client.post(
                "/api/cloud/project-invites", json={"role": "capturer"}
            )
            updated = self.client.patch(
                "/api/cloud/project-members/77", json={"role": "viewer"}
            )
            removed = self.client.delete("/api/cloud/project-members/77")
            revoked = self.client.delete("/api/cloud/project-invites/9")

        self.assertEqual(members.status_code, 200, members.get_json())
        self.assertEqual(members.get_json()["members"][0]["role"], "owner")
        self.assertEqual(invite.status_code, 200, invite.get_json())
        self.assertEqual(updated.status_code, 200, updated.get_json())
        self.assertEqual(removed.status_code, 200, removed.get_json())
        self.assertEqual(revoked.status_code, 200, revoked.get_json())
        get_members.assert_called_once_with("owner-local-project", "ARCHIVE-PC")
        create_invite.assert_called_once_with(
            "owner-local-project", "capturer", "ARCHIVE-PC"
        )
        update_member.assert_called_once_with(
            "77", "owner-local-project", "viewer", "ARCHIVE-PC"
        )
        remove_member.assert_called_once_with("77", "owner-local-project", "ARCHIVE-PC")
        revoke_invite.assert_called_once_with("9", "owner-local-project", "ARCHIVE-PC")

        with patch.object(
            self.desktop,
            "current_project_cloud_ownership",
            return_value={"state": "mismatch"},
        ), patch.object(
            self.desktop.cloud_client, "desktop_project_members", new=MagicMock()
        ) as blocked_call:
            blocked = self.client.get("/api/cloud/project-members")
        self.assertEqual(blocked.status_code, 409)
        blocked_call.assert_not_called()


if __name__ == "__main__":
    unittest.main()
